from openpyxl import load_workbook


def excelhandling(filepath):
    with open(filepath) as data:
        workbook = load_workbook("testdata\\sample_creds.xlsx")
        sheet = workbook["Sheet2"]
        values = []
        for i in sheet.iter_rows(min_row=1,values_only=True):
            values.append(i)
        return values